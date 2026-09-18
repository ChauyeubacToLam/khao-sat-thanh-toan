"""Xuất bảng hỏi ra file Word theo bố cục của phiếu mẫu:
thư ngỏ, thông tin chung, rồi mỗi khái niệm một mục có định nghĩa và bảng Code | Item | 1..5.

Nội dung đọc thẳng từ survey.js để bản Word và bản web không bao giờ lệch nhau.

    python make_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

HERE = Path(__file__).parent
OUT = HERE.parent / "Phieu khao sat - Thanh toan tu phuc vu.docx"

LIKERT = ["Hoàn toàn không đồng ý", "Không đồng ý", "Trung lập",
          "Đồng ý", "Hoàn toàn đồng ý"]


def parse_survey_js():
    src = (HERE / "survey.js").read_text(encoding="utf-8")

    blocks = []
    for m in re.finditer(
            r'\{\s*id:\s*"(\w+)",\s*title:\s*"([^"]+)",\s*def:\s*"([^"]+)",\s*'
            r'src:\s*"([^"]+)",\s*items:\s*\[(.*?)\],\s*\},', src, re.S):
        bid, title, define, source, body = m.groups()
        items = re.findall(
            r'\["(\w+)",\s*"((?:[^"\\]|\\.)*)"(,\s*\{\s*reverse:\s*true\s*\})?\]', body)
        blocks.append(dict(id=bid, title=title, define=define, src=source,
                           items=[(c, t.replace('\\"', '"'), bool(r)) for c, t, r in items]))

    qc = [dict(code=c, block=b, text=t) for c, b, t in re.findall(
        r'\{\s*code:\s*"(\w+)",\s*block:\s*"(\w+)",[^}]*?text:\s*"([^"]+)"', src, re.S)]

    sat = re.findall(r'\["(SAT\d)",\s*"([^"]+)",\s*"([^"]+)"\]', src)

    demo = []
    demo_src = re.search(r'const DEMO = \[(.*?)\n\];', src, re.S).group(1)
    for m in re.finditer(r'\{\s*code:\s*"(\w+)",\s*label:\s*"([^"]+)",\s*opts:\s*\[(.*?)\]\s*\}',
                         demo_src, re.S):
        demo.append((m.group(1), m.group(2), re.findall(r'"([^"]+)"', m.group(3))))

    err = []
    err_src = re.search(r'const ERROR_BLOCK = \[(.*?)\n\];', src, re.S).group(1)
    for m in re.finditer(r'\{\s*code:\s*"(\w+)",\s*label:\s*"([^"]+)",(.*?)opts:\s*\[(.*?)\]\s*\}',
                         err_src, re.S):
        err.append((m.group(1), m.group(2), "multi: true" in m.group(3),
                    re.findall(r'"([^"]+)"', m.group(4))))

    op = re.findall(r'\{\s*code:\s*"(open_\w+)",\s*label:\s*"([^"]+)"\s*\}', src)
    return blocks, qc, sat, demo, err, op


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)


def setup(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.25
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)


def h(doc, text, size=13, before=14, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    return p


def para(doc, text, italic=False, size=12, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def likert_table(doc, rows, header):
    t = doc.add_table(rows=1, cols=7)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ["Mã", header] + ["%d\n%s" % (i + 1, LIKERT[i]) for i in range(5)]
    for c, txt in zip(t.rows[0].cells, labels):
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        r.bold = True
        r.font.size = Pt(9)
        shade(c, "E8F1EE")
    for code, text in rows:
        cells = t.add_row().cells
        cells[0].text = ""
        r = cells[0].paragraphs[0].add_run(code)
        r.bold = True
        r.font.size = Pt(10)
        cells[1].text = ""
        cells[1].paragraphs[0].add_run(text).font.size = Pt(11)
        for i in range(5):
            cells[2 + i].text = ""
            p = cells[2 + i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(str(i + 1)).font.size = Pt(11)
    widths = [Cm(1.5), Cm(7.4), Cm(1.4), Cm(1.4), Cm(1.4), Cm(1.4), Cm(1.4)]
    for row in t.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    return t


def choice_block(doc, idx, label, opts, multi=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("%d. %s" % (idx, label))
    r.bold = True
    r.font.size = Pt(11.5)
    if multi:
        p.add_run("  (có thể chọn nhiều đáp án)").italic = True
    for o in opts:
        q = doc.add_paragraph()
        q.paragraph_format.left_indent = Cm(0.8)
        q.paragraph_format.space_after = Pt(2)
        q.add_run("☐  " + o).font.size = Pt(11)


def build():
    blocks, qc, sat, demo, err, op = parse_survey_js()
    qc_by_block = {}
    for q in qc:
        qc_by_block.setdefault(q["block"], []).append(q)

    doc = Document()
    setup(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PHIẾU KHẢO SÁT")
    r.bold = True
    r.font.size = Pt(16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SỰ HÀI LÒNG CỦA KHÁCH HÀNG VỚI CÁC GIẢI PHÁP\n"
                  "THANH TOÁN TỰ PHỤC VỤ TẠI CỬA HÀNG BÁN LẺ")
    r.bold = True
    r.font.size = Pt(14)
    para(doc, "Trường Đại học Ngoại thương", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph()
    para(doc, "Thân gửi Anh/Chị và các bạn,")
    para(doc, "Anh/Chị đã bao giờ tự quét mã món hàng rồi tự trả tiền ở siêu thị, tự gọi món trên "
              "màn hình cảm ứng ở quán ăn nhanh, hay quét mã QR để thanh toán ngay tại quầy chưa?")
    para(doc, "Anh/Chị thấy việc đó nhanh, chủ động và tiện lợi, hay đôi lúc lại thấp thỏm lo máy "
              "quét sót món, tính tiền trùng, trừ sai số tiền mà không biết gọi ai để xử lý?")
    para(doc, "Trong bối cảnh các chuỗi bán lẻ tại Việt Nam triển khai ngày càng nhiều quầy thanh "
              "toán tự phục vụ, việc hiểu đúng cảm nhận của khách hàng trở thành vấn đề quan trọng. "
              "Công nghệ giúp giao dịch nhanh hơn, nhưng đồng thời chuyển phần công việc kiểm tra và "
              "xử lý sai sót sang phía khách hàng. Nhóm nghiên cứu thực hiện khảo sát này nhằm đo "
              "lường cán cân giữa hai mặt đó, từ đó đề xuất giải pháp giúp doanh nghiệp bán lẻ nâng "
              "cao trải nghiệm mua sắm và giảm thiểu sai sót trong thanh toán.")
    para(doc, "Khảo sát dự kiến mất khoảng 7 tới 9 phút để hoàn thành. Mọi thông tin Anh/Chị cung "
              "cấp sẽ được bảo mật tuyệt đối và chỉ phục vụ mục đích nghiên cứu khoa học. Không có "
              "câu trả lời đúng hay sai, nhóm chỉ cần cảm nhận thật của Anh/Chị.")
    para(doc, "Nhóm nghiên cứu xin chân thành cảm ơn.")

    h(doc, "CÂU HỎI SÀNG LỌC")
    choice_block(doc, 1,
                 "Trong 3 THÁNG QUA, Anh/Chị đã từng sử dụng máy hoặc quầy thanh toán tự phục vụ "
                 "tại cửa hàng bán lẻ chưa? (quầy tự quét tự trả tiền, quầy bán tự động, kiosk tự "
                 "gọi món, quét mã QR tại quầy)",
                 ["Rồi, tôi đã từng sử dụng",
                  "Chưa, tôi chưa từng sử dụng (dừng khảo sát tại đây)"])

    h(doc, "PHẦN I. THÔNG TIN CHUNG")
    para(doc, "Thông tin dưới đây chỉ dùng để phân nhóm khi phân tích, không định danh người trả lời.",
         italic=True, size=11)
    for i, (code, label, opts) in enumerate(demo, 1):
        choice_block(doc, i, label, opts)

    para(doc, "")
    h(doc, "PHẦN II. CÁC CÂU HỎI VỀ CẢM NHẬN")
    para(doc, "Anh/Chị vui lòng đánh giá mức độ đồng ý với từng phát biểu theo thang điểm từ 1 tới "
              "5, trong đó: 1 = Hoàn toàn không đồng ý, 2 = Không đồng ý, 3 = Trung lập, "
              "4 = Đồng ý, 5 = Hoàn toàn đồng ý.", size=11)

    roman = ["II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI"]
    for n, b in enumerate(blocks):
        h(doc, "%s. %s" % (roman[n], b["title"]), size=12.5)
        para(doc, b["define"], size=11)
        rows = [(c, t + (" (câu đảo chiều)" if rev else "")) for c, t, rev in b["items"]]
        for q in qc_by_block.get(b["id"], []):
            rows.append((q["code"], q["text"] + "  [câu kiểm soát chất lượng]"))
        likert_table(doc, rows, b["title"])
        para(doc, "Nguồn thang đo: " + b["src"], italic=True, size=9.5)

    h(doc, "%s. Mức độ hài lòng tổng thể" % roman[len(blocks)], size=12.5)
    para(doc, "Nhìn lại toàn bộ trải nghiệm khi thanh toán tại máy tự phục vụ, Anh/Chị vui lòng "
              "chọn vị trí gần với cảm nhận của mình nhất trên thang từ 1 tới 5.", size=11)
    t = doc.add_table(rows=1, cols=8)
    t.style = "Table Grid"
    for c, txt in zip(t.rows[0].cells, ["Mã", "Trải nghiệm đó là", "1", "2", "3", "4", "5", ""]):
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(txt)
        r.bold = True
        r.font.size = Pt(9)
        shade(c, "E8F1EE")
    for code, lo, hi in sat:
        cells = t.add_row().cells
        cells[0].text = ""
        cells[0].paragraphs[0].add_run(code).bold = True
        cells[1].text = ""
        cells[1].paragraphs[0].add_run(lo).font.size = Pt(10.5)
        for i in range(5):
            cells[2 + i].text = ""
            p = cells[2 + i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(str(i + 1)).font.size = Pt(11)
        cells[7].text = ""
        cells[7].paragraphs[0].add_run(hi).font.size = Pt(10.5)
    for row in t.rows:
        for cell, w in zip(row.cells, [Cm(1.3), Cm(3.6), Cm(1.1), Cm(1.1), Cm(1.1),
                                       Cm(1.1), Cm(1.1), Cm(3.6)]):
            cell.width = w
    para(doc, "Nguồn thang đo: Bhattacherjee (2001), Satisfaction (semantic differential)",
         italic=True, size=9.5)

    para(doc, "")
    h(doc, "PHẦN III. TRẢI NGHIỆM SỰ CỐ THANH TOÁN")
    para(doc, "Đây là phần quan trọng nhất với nghiên cứu, mong Anh/Chị trả lời thật.",
         italic=True, size=11)
    for i, (code, label, multi, opts) in enumerate(err, 1):
        choice_block(doc, i, label, opts, multi)
    for i, (code, label) in enumerate(op, len(err) + 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run("%d. %s" % (i, label))
        r.bold = True
        r.font.size = Pt(11.5)
        for _ in range(2):
            d = doc.add_paragraph("." * 96)
            d.paragraph_format.left_indent = Cm(0.8)

    para(doc, "")
    para(doc, "Nhóm nghiên cứu xin chân thành cảm ơn Anh/Chị đã dành thời gian hoàn thành khảo sát. "
              "Chúc Anh/Chị một ngày tốt lành.", align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_page_break()
    h(doc, "PHỤ LỤC (dành riêng cho nhóm nghiên cứu, không in cho người trả lời)", size=12)
    para(doc, "Các câu kiểm soát chất lượng dưới đây được trộn ngẫu nhiên vào từng khối trên bản "
              "trực tuyến. Phiếu vi phạm sẽ bị đánh dấu để rà soát hoặc loại khỏi phân tích.", size=11)
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for c, txt in zip(t.rows[0].cells, ["Mã", "Loại", "Nằm trong khối", "Tiêu chí hợp lệ"]):
        c.text = ""
        r = c.paragraphs[0].add_run(txt)
        r.bold = True
        r.font.size = Pt(10)
        shade(c, "E8F1EE")
    rules = [
        ("AC1", "Câu bẫy trực tiếp", "SEC", "Phải chọn đúng mức 2"),
        ("AC2", "Câu bẫy trực tiếp", "PRK", "Phải chọn đúng mức 4"),
        ("BOG", "Câu vô lý", "TAX", "Phải chọn mức 1 hoặc 2"),
        ("SYN1, SYN2", "Cặp đồng nghĩa", "REL và WOM", "Hai câu chênh nhau dưới 3 mức"),
        ("ANT1, ANT2", "Cặp trái nghĩa", "EFF và CI", "Hai câu chênh nhau từ 2 mức trở lên"),
        ("CI3", "Item đảo chiều", "CI", "Phải ngược hướng với CI1"),
    ]
    for row in rules:
        for cell, txt in zip(t.add_row().cells, row):
            cell.text = ""
            cell.paragraphs[0].add_run(txt).font.size = Pt(10)

    n_items = sum(len(b["items"]) for b in blocks) + len(sat)
    para(doc, "")
    para(doc, "Tổng số item thang đo: %d. Tổng số câu kiểm soát: %d. Thang đo Likert 5 mức. "
              "Các thang gốc đều dùng 7 mức và đã được chuyển về 5 mức, cần khai báo điều này "
              "trong phần Phương pháp nghiên cứu." % (n_items, len(qc)), size=10.5)

    doc.save(OUT)
    print("Da luu:", OUT)
    print("So khoi:", len(blocks), "| item thang do:", n_items, "| cau kiem soat:", len(qc))


if __name__ == "__main__":
    build()
